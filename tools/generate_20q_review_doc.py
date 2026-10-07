"""Gera DOCX com 20 perguntas respondidas pelo fluxo real da POC.

Executa, para cada pergunta:
1. resposta RAG com openai/gpt-4o-mini;
2. resposta baseline com openai/gpt-4o-mini;
3. avaliação automática com 3 avaliadores baratos;
4. persistência no histórico da aplicação;
5. montagem de um DOCX para revisão humana.
"""

from __future__ import annotations

import re
import time
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from src.history_store import save_interaction
from src.quality_judge import evaluate_answers_with_judges
from src.rag_chain import answer_question


RESPONSE_MODEL = "openai/gpt-4o-mini"
RESPONSE_MODEL_LABEL = "Pago · GPT-4o mini · parâmetros não divulgados · ctx 128K"
JUDGE_MODELS = [
    "google/gemini-2.5-flash-lite",
    "deepseek/deepseek-v4-flash",
    "amazon/nova-lite-v1",
]
JUDGE_LABELS = {
    "google/gemini-2.5-flash-lite": "Gemini 2.5 Flash Lite",
    "deepseek/deepseek-v4-flash": "DeepSeek V4 Flash",
    "amazon/nova-lite-v1": "Amazon Nova Lite",
}
OUTPUT = Path("/app/avaliacao_20_perguntas_cdc_gpt4o_mini.docx")


QUESTIONS = [
    {
        "tema": "Direito de arrependimento em compra online",
        "pergunta": (
            "Comprei um produto pela internet, ele chegou na minha casa, mas eu "
            "me arrependi da compra. Em quais situações posso desistir, qual é "
            "o prazo para fazer isso e o que o fornecedor precisa devolver?"
        ),
    },
    {
        "tema": "Produto com defeito e prazo para solução",
        "pergunta": (
            "Comprei um produto novo e ele apresentou defeito poucos dias depois. "
            "A loja disse que vai mandar para assistência, mas eu quero entender "
            "quais são meus direitos se o problema não for resolvido no prazo legal."
        ),
    },
    {
        "tema": "Cobrança indevida e devolução em dobro",
        "pergunta": (
            "Recebi uma cobrança que considero indevida e acabei pagando para evitar "
            "restrição no meu nome. O CDC permite pedir devolução em dobro? Em quais "
            "casos isso se aplica?"
        ),
    },
    {
        "tema": "Nome negativado após dívida paga",
        "pergunta": (
            "Paguei uma dívida que estava em atraso, mas meu nome continua aparecendo "
            "como negativado no Serasa ou SPC. Qual é o prazo para retirar a restrição "
            "e o que posso fazer se isso não acontecer?"
        ),
    },
    {
        "tema": "Publicidade enganosa",
        "pergunta": (
            "Vi uma oferta anunciada com determinadas características e preço, mas "
            "quando tentei comprar a loja informou condições diferentes. Como o CDC "
            "trata publicidade enganosa e quais direitos tenho?"
        ),
    },
    {
        "tema": "Venda casada",
        "pergunta": (
            "Uma empresa disse que só vende um produto se eu contratar também outro "
            "serviço junto. Isso é permitido pelo Código de Defesa do Consumidor? "
            "O que caracteriza venda casada?"
        ),
    },
    {
        "tema": "Cartão de crédito não solicitado",
        "pergunta": (
            "Recebi em casa um cartão de crédito que eu nunca solicitei. A empresa "
            "pode enviar cartão sem pedido do consumidor? Isso gera alguma consequência "
            "pelo CDC?"
        ),
    },
    {
        "tema": "Fraude bancária e transação desconhecida",
        "pergunta": (
            "Apareceu uma movimentação bancária ou transação no cartão que eu não "
            "reconheço, possivelmente causada por golpe. O CDC se aplica aos bancos "
            "e qual pode ser a responsabilidade da instituição financeira?"
        ),
    },
    {
        "tema": "Cancelamento de serviço recorrente",
        "pergunta": (
            "Tenho uma assinatura de serviço recorrente e quero cancelar, mas a empresa "
            "dificulta o cancelamento e continua cobrando. Quais direitos o consumidor "
            "tem nessa situação?"
        ),
    },
    {
        "tema": "Garantia legal e garantia contratual",
        "pergunta": (
            "Qual é a diferença entre garantia legal e garantia contratual? Se a loja "
            "diz que a garantia acabou, ainda posso reclamar com base no CDC?"
        ),
    },
    {
        "tema": "Vício oculto descoberto depois",
        "pergunta": (
            "Comprei um produto e só descobri o defeito depois de algum tempo de uso, "
            "porque o problema não era aparente. Como funciona o prazo para reclamar "
            "de vício oculto no CDC?"
        ),
    },
    {
        "tema": "Atraso na entrega de compra online",
        "pergunta": (
            "Comprei pela internet com prazo de entrega definido, mas o produto atrasou "
            "e a loja não dá solução clara. Quais opções o consumidor tem quando o "
            "fornecedor não cumpre o prazo de entrega?"
        ),
    },
    {
        "tema": "Preço diferente entre anúncio e caixa",
        "pergunta": (
            "O produto estava anunciado por um preço, mas no caixa ou no site apareceu "
            "outro valor maior. O fornecedor é obrigado a cumprir o preço anunciado?"
        ),
    },
    {
        "tema": "Cláusula abusiva em contrato",
        "pergunta": (
            "Assinei um contrato de consumo com multa muito alta e cláusulas que parecem "
            "colocar o consumidor em desvantagem exagerada. Como o CDC trata cláusulas "
            "abusivas?"
        ),
    },
    {
        "tema": "Serviço mal prestado",
        "pergunta": (
            "Contratei um serviço, mas ele foi prestado com falhas e não resolveu o "
            "problema prometido. O que o consumidor pode exigir quando há defeito ou "
            "má prestação de serviço?"
        ),
    },
    {
        "tema": "Plano de saúde e negativa de cobertura",
        "pergunta": (
            "Meu plano de saúde negou cobertura para um procedimento indicado pelo médico. "
            "Em termos de relação de consumo, o CDC pode ser aplicado a planos de saúde?"
        ),
    },
    {
        "tema": "Construtora, imóvel e devolução de valores",
        "pergunta": (
            "Comprei um imóvel na planta e quero desistir do contrato. A construtora pode "
            "reter todo o valor pago? Como o CDC costuma proteger o consumidor nesses "
            "contratos?"
        ),
    },
    {
        "tema": "Produto ou serviço perigoso e recall",
        "pergunta": (
            "Descobri que um produto comprado pode apresentar risco à saúde ou segurança "
            "do consumidor. Quais deveres de informação e prevenção o fornecedor tem "
            "pelo CDC?"
        ),
    },
    {
        "tema": "Orçamento prévio e cobrança em assistência técnica",
        "pergunta": (
            "Levei um produto para conserto e a assistência começou o serviço sem orçamento "
            "claro, depois cobrou um valor alto. O CDC exige orçamento prévio ou autorização "
            "do consumidor?"
        ),
    },
    {
        "tema": "Atendimento, informação clara e protocolo",
        "pergunta": (
            "Estou tentando resolver um problema com uma empresa, mas recebo informações "
            "confusas, sem protocolo e sem resposta objetiva. Quais direitos básicos do "
            "consumidor ajudam nesse tipo de situação?"
        ),
    },
]


def retry(fn, attempts: int = 3, delay: int = 3):
    last = None
    for attempt in range(1, attempts + 1):
        try:
            return fn()
        except Exception as exc:  # pragma: no cover - operação externa
            last = exc
            print(
                f"    tentativa {attempt}/{attempts} falhou: {str(exc)[:180]}",
                flush=True,
            )
            if attempt < attempts:
                time.sleep(delay * attempt)
    raise last


def clean_md(text: str) -> str:
    text = text or ""
    text = text.replace("\\\n", "\n")
    text = re.sub(r"\*\*(.*?)\*\*", r"\1", text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = text.replace("### ", "").replace("## ", "").replace("# ", "")
    text = text.replace("\\", "")
    return text.strip()


def doc_refs(result: dict) -> str:
    refs = []
    for doc in result.get("documents", [])[:3]:
        refs.append(doc.metadata.get("artigo", "CDC"))
    for doc in result.get("history_documents", [])[:2]:
        refs.append(doc.metadata.get("referencia", "Histórico CDC"))
    for doc in result.get("jurisprudence_documents", [])[:2]:
        refs.append(doc.metadata.get("referencia", "STJ"))

    seen = []
    for ref in refs:
        if ref and ref not in seen:
            seen.append(ref)
    return ", ".join(seen) if seen else "Sem fontes recuperadas"


def run_case(item: dict, index: int) -> dict:
    question = item["pergunta"]
    print(f"[{index:02d}/20] {item['tema']}", flush=True)
    rag = retry(
        lambda: answer_question(question, use_rag=True, model_name=RESPONSE_MODEL)
    )
    time.sleep(0.8)
    baseline = retry(
        lambda: answer_question(question, use_rag=False, model_name=RESPONSE_MODEL)
    )
    time.sleep(0.8)
    quality = evaluate_answers_with_judges(
        question=question,
        rag_result=rag,
        baseline_result=baseline,
        judge_models=JUDGE_MODELS,
    )
    for result in quality.get("judge_results", []):
        model_id = result.get("effective_judge_model") or result.get("judge_model")
        result["judge_model_display"] = JUDGE_LABELS.get(model_id, model_id)
    quality["judge_models_display"] = [
        JUDGE_LABELS.get(model, model) for model in JUDGE_MODELS
    ]
    history_id = save_interaction(
        question=question,
        mode="Comparação RAG x Baseline",
        compare_mode=True,
        response_model=RESPONSE_MODEL,
        response_model_label=RESPONSE_MODEL_LABEL,
        api_key_label="Chave usada pelo ambiente",
        rag_result=rag,
        baseline_result=baseline,
        ai_quality=quality,
    )
    return {
        "index": index,
        "tema": item["tema"],
        "pergunta": question,
        "rag": rag,
        "baseline": baseline,
        "quality": quality,
        "history_id": history_id,
        "fontes": doc_refs(rag),
    }


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(str(text))
    run.bold = bold
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(8.5)
        paragraph.paragraph_format.space_after = Pt(0)


def add_kv_table(doc: Document, rows: list[tuple[str, str]]):
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    table.autofit = False
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].width = Inches(1.55)
        cells[1].width = Inches(4.85)
        set_cell_shading(cells[0], "F2F4F7")
        set_cell_text(cells[0], label, bold=True)
        set_cell_text(cells[1], value)
        cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cells[1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    doc.add_paragraph()
    return table


def add_answer_block(doc: Document, title: str, text: str) -> None:
    doc.add_heading(title, level=3)
    cleaned = clean_md(text)
    if not cleaned:
        doc.add_paragraph("Sem resposta registrada.")
        return

    for block in re.split(r"\n\s*\n", cleaned):
        block = block.strip()
        if not block:
            continue
        if block.startswith(("- ", "* ")):
            for line in block.splitlines():
                line = line.strip().lstrip("-* ").strip()
                if line:
                    doc.add_paragraph(line, style="List Bullet")
        else:
            paragraph = doc.add_paragraph(block)
            paragraph.paragraph_format.space_after = Pt(6)


def add_quality_table(doc: Document, quality: dict) -> None:
    rag_score = quality.get("rag", {}).get("score", "-")
    baseline_score = quality.get("baseline", {}).get("score", "-")
    winner = quality.get("winner", "Indefinido")
    summary = quality.get("summary", "")
    add_kv_table(
        doc,
        [
            ("Nota RAG", f"{rag_score}/5"),
            ("Nota Baseline", f"{baseline_score}/5"),
            ("Melhor resposta", winner),
            ("Resumo", summary),
        ],
    )

    judges = quality.get("judge_results", [])
    if judges:
        paragraph = doc.add_paragraph("Notas individuais dos avaliadores")
        paragraph.runs[0].bold = True
        table = doc.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        headers = ["Avaliador", "RAG", "Baseline", "Melhor"]
        for i, header in enumerate(headers):
            set_cell_shading(table.rows[0].cells[i], "E8EEF5")
            set_cell_text(table.rows[0].cells[i], header, bold=True)
        for result in judges:
            row = table.add_row().cells
            set_cell_text(
                row[0],
                result.get("judge_model_display") or result.get("judge_model"),
            )
            set_cell_text(row[1], result.get("rag", {}).get("score", "-"))
            set_cell_text(row[2], result.get("baseline", {}).get("score", "-"))
            set_cell_text(row[3], result.get("winner", "-"))
        doc.add_paragraph()


def add_human_review_form(doc: Document) -> None:
    doc.add_heading("Campo para avaliação humana", level=3)
    table = doc.add_table(rows=3, cols=2)
    table.style = "Table Grid"
    rows = [
        ("Nota humana (0 a 5)", "____ / 5"),
        ("Avaliador humano", "____________________________________________"),
        ("Comentário técnico", "\n\n\n"),
    ]
    for idx, (label, value) in enumerate(rows):
        cells = table.rows[idx].cells
        cells[0].width = Inches(1.8)
        cells[1].width = Inches(4.6)
        set_cell_shading(cells[0], "F2F4F7")
        set_cell_text(cells[0], label, bold=True)
        set_cell_text(cells[1], value)
    doc.add_paragraph()


def setup_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    title = styles["Title"]
    title.font.name = "Calibri"
    title.font.size = Pt(22)
    title.font.bold = True
    title.font.color.rgb = RGBColor(11, 37, 69)
    title.paragraph_format.space_after = Pt(6)

    for style_name, size, color in [
        ("Heading 1", 16, RGBColor(46, 116, 181)),
        ("Heading 2", 13, RGBColor(46, 116, 181)),
        ("Heading 3", 12, RGBColor(31, 77, 120)),
    ]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(8)
        style.paragraph_format.space_after = Pt(4)


def create_doc(results: list[dict]) -> None:
    doc = Document()
    setup_styles(doc)

    title = doc.add_paragraph(style="Title")
    title.add_run("Relatório de avaliação - 20 perguntas sobre CDC")
    subtitle = doc.add_paragraph()
    subtitle.add_run("Chatbot RAG - Código de Defesa do Consumidor").bold = True
    doc.add_paragraph(
        f"Documento gerado em {datetime.now().strftime('%d/%m/%Y %H:%M')}, "
        "usando o fluxo da aplicação: resposta RAG, resposta baseline e "
        "avaliação automática de qualidade."
    )
    add_kv_table(
        doc,
        [
            ("Modelo de resposta", RESPONSE_MODEL_LABEL),
            (
                "Avaliadores automáticos",
                ", ".join(JUDGE_LABELS[model] for model in JUDGE_MODELS),
            ),
            (
                "Escala de qualidade",
                "0 a 5, onde 5 indica resposta mais completa, fundamentada, "
                "clara e com menor risco de alucinação.",
            ),
            (
                "Objetivo do documento",
                "Permitir revisão humana das respostas geradas pela POC e "
                "coleta de feedback técnico da área.",
            ),
        ],
    )

    doc.add_heading("Resumo executivo das avaliações", level=1)
    table = doc.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    headers = ["#", "Tema", "Nota RAG", "Nota Baseline", "Melhor", "Histórico"]
    for i, header in enumerate(headers):
        set_cell_shading(table.rows[0].cells[i], "E8EEF5")
        set_cell_text(table.rows[0].cells[i], header, bold=True)
    for result in results:
        quality = result["quality"]
        row = table.add_row().cells
        set_cell_text(row[0], result["index"])
        set_cell_text(row[1], result["tema"])
        set_cell_text(row[2], quality.get("rag", {}).get("score", "-"))
        set_cell_text(row[3], quality.get("baseline", {}).get("score", "-"))
        set_cell_text(row[4], quality.get("winner", "-"))
        set_cell_text(row[5], f"#{result.get('history_id')}")
    doc.add_page_break()

    for result in results:
        doc.add_heading(f"{result['index']}. {result['tema']}", level=1)
        add_kv_table(
            doc,
            [
                ("Pergunta", result["pergunta"]),
                ("Fontes RAG recuperadas", result.get("fontes", "Sem fontes")),
                ("ID no histórico da aplicação", f"#{result.get('history_id')}"),
            ],
        )
        add_answer_block(doc, "Resposta com RAG", result["rag"].get("answer", ""))
        add_answer_block(
            doc,
            "Resposta baseline sem recuperação",
            result["baseline"].get("answer", ""),
        )
        doc.add_heading("Avaliação automática de qualidade", level=2)
        add_quality_table(doc, result["quality"])
        add_human_review_form(doc)
        if result["index"] != len(results):
            doc.add_page_break()

    doc.core_properties.title = "Relatório de avaliação - 20 perguntas CDC"
    doc.core_properties.subject = "TCC NLP - RAG CDC"
    doc.core_properties.author = "Equipe TCC"
    doc.save(OUTPUT)


def error_result(item: dict, index: int, exc: Exception) -> dict:
    return {
        "index": index,
        "tema": item["tema"],
        "pergunta": item["pergunta"],
        "rag": {
            "answer": f"ERRO: {exc}",
            "documents": [],
            "history_documents": [],
            "jurisprudence_documents": [],
        },
        "baseline": {"answer": "Não gerado por erro anterior."},
        "quality": {
            "rag": {"score": 0},
            "baseline": {"score": 0},
            "winner": "Indefinido",
            "summary": f"Erro técnico: {exc}",
            "judge_results": [],
        },
        "history_id": "não salvo",
        "fontes": "não recuperadas por erro",
    }


def main() -> None:
    results = []
    for index, item in enumerate(QUESTIONS, start=1):
        try:
            results.append(run_case(item, index))
        except Exception as exc:  # pragma: no cover - operação externa
            print(f"ERRO na pergunta {index}: {exc}", flush=True)
            results.append(error_result(item, index, exc))
        time.sleep(1)

    create_doc(results)
    print(f"DOCX gerado: {OUTPUT}", flush=True)


if __name__ == "__main__":
    main()
